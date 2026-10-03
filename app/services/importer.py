from pathlib import Path
from datetime import date,datetime,time,timedelta
from openpyxl import load_workbook
from sqlalchemy import select
from app.database import db
from app.database.models import Staff,Admin,Lesson,Requirement,Score,Schedule,Attendance,Setting
BASE=date(1899,12,30)
def as_date(v):
    if v in (None,""):return None
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    if isinstance(v,(int,float)):return BASE+timedelta(days=int(v))
    for f in ("%m/%d/%Y","%Y-%m-%d","%d-%b-%Y"):
        try:return datetime.strptime(str(v).strip(),f).date()
        except ValueError:pass
    return None
def as_time(v):
    if v in (None,""):return None
    if isinstance(v,time):return v
    if isinstance(v,datetime):return v.time().replace(second=0,microsecond=0)
    if isinstance(v,(int,float)):
        sec=round(float(v)*86400);return time((sec//3600)%24,(sec%3600)//60)
    for f in ("%H:%M","%I:%M %p"):
        try:return datetime.strptime(str(v).strip(),f).time()
        except ValueError:pass
    return None
def yes(v):return str(v).strip().lower() in {"yes","y","true","1"}
def rows(path,sheet):
    wb=load_workbook(path,data_only=True,read_only=True);ws=wb[sheet];headers=[str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    for values in ws.iter_rows(min_row=2,values_only=True):
        if any(v not in (None,"") for v in values):yield dict(zip(headers,values))
    wb.close()
def upsert(s,model,key,match_value,**values):
    obj=s.scalar(select(model).where(getattr(model,key)==match_value))
    if obj:
        for k,v in values.items():setattr(obj,k,v)
    else:s.add(model(**{key:match_value},**values))
def import_all(folder):
    p=Path(folder);s=db.SessionLocal();summary={}
    try:
        if (p/"01_Staff_Master.xlsx").exists():
            for r in rows(p/"01_Staff_Master.xlsx","Staff_Master"):
                sid=str(r["Staff ID"]).strip();chat=r.get("Telegram Chat ID");upsert(s,Staff,"staff_id",sid,name=str(r["Staff Name"]).strip(),department=r.get("Department"),section=r.get("Section"),position=r.get("Position"),supervisor_id=str(r.get("Supervisor ID") or "").strip() or None,supervisor_name=r.get("Supervisor Name"),is_supervisor=yes(r.get("Is Supervisor")),active=yes(r.get("Active")),shift=r.get("Shift"),telegram_chat_id=int(chat) if chat not in (None,"") else None,join_date=as_date(r.get("Join Date")),location=r.get("Location"))
        if (p/"02_Lesson_Master.xlsx").exists():
            for r in rows(p/"02_Lesson_Master.xlsx","Lesson_Master"):
                code=str(r["Lesson Code"]).strip();upsert(s,Lesson,"code",code,category=str(r["Category"]),name=str(r["Lesson Name"]),pass_score=int(r.get("Pass Score") or 50),validity_months=int(r.get("Validity Months") or 0),active=yes(r.get("Active")),description=r.get("Description"))
        if (p/"03_Lesson_Requirements.xlsx").exists() and s.query(Requirement).count()==0:
            for r in rows(p/"03_Lesson_Requirements.xlsx","Lesson_Requirements"):s.add(Requirement(department=str(r.get("Department") or "All"),section=str(r.get("Section") or "All"),position=str(r.get("Position") or "All"),category=str(r["Category"]),lesson_code=str(r["Lesson Code"]),required=yes(r.get("Required")),effective_date=as_date(r.get("Effective Date")),active=yes(r.get("Active"))))
        if (p/"05_Admin_List.xlsx").exists():
            for r in rows(p/"05_Admin_List.xlsx","Admin_List"):
                chat=r.get("Telegram Chat ID");upsert(s,Admin,"staff_id",str(r["Staff ID"]),name=str(r["Admin Name"]),role=str(r["Admin Role"]),telegram_chat_id=int(chat) if chat not in (None,"") else None,active=yes(r.get("Active")))
        if (p/"04_Existing_Scores.xlsx").exists() and s.query(Score).count()==0:
            for r in rows(p/"04_Existing_Scores.xlsx","Scores"):
                dt=as_date(r.get("Assessment Date"))
                if dt:s.add(Score(staff_id=str(r["Staff ID"]),category=str(r["Category"]),lesson_code=str(r["Lesson Code"]),score=float(r["Score"]),assessment_date=dt,training_date=as_date(r.get("Training Date")),remark=r.get("Remark")))
        if (p/"07_Current_Schedules.xlsx").exists() and s.query(Schedule).count()==0:
            location=None
            for r in rows(p/"07_Current_Schedules.xlsx","Schedules"):
                location=r.get("Location") or location;dt=as_date(r.get("Date"));start=as_time(r.get("Start Time"));end=as_time(r.get("End Time"))
                if dt and start and end and location:s.add(Schedule(category=str(r["Category"]),lesson_code=str(r["Lesson Code"]),session_no=int(r["Session No."]),session_date=dt,start_time=start,end_time=end,location=location,status=str(r.get("Status") or "Active"),remark=r.get("Remark")))
        if (p/"08_Attendance_History.xlsx").exists() and s.query(Attendance).count()==0:
            for r in rows(p/"08_Attendance_History.xlsx","Attendance_History"):
                dt=as_date(r.get("Session Date"))
                if dt:s.add(Attendance(staff_id=str(r["Staff ID"]),category=str(r["Category"]),lesson_code=str(r["Lesson Code"]),session_date=dt,start_time=as_time(r.get("Start Time")),attendance=str(r.get("Attendance") or "Present"),location=r.get("Location"),remark=r.get("Remark")))
        if (p/"11_System_Settings.xlsx").exists():
            for r in rows(p/"11_System_Settings.xlsx","Settings"):upsert(s,Setting,"key",str(r["Setting"]),value=str(r.get("Value") or ""),notes=r.get("Notes"))
        s.commit();summary={"staff":s.query(Staff).count(),"lessons":s.query(Lesson).count(),"scores":s.query(Score).count(),"schedules":s.query(Schedule).count()};return summary
    except Exception:s.rollback();raise
    finally:s.close()
def ensure_bootstrap(staff_id,name):
    s=db.SessionLocal()
    try:
        if not s.scalar(select(Staff).where(Staff.staff_id==staff_id)):s.add(Staff(staff_id=staff_id,name=name,department="Training",section="TPM",position="Super Admin",is_supervisor=True,active=True))
        if not s.scalar(select(Admin).where(Admin.staff_id==staff_id)):s.add(Admin(staff_id=staff_id,name=name,role="Super Admin",active=True))
        s.commit()
    finally:s.close()
